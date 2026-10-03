import io
import fitz  # PyMuPDF
import docx
import pytest
from fastapi.testclient import TestClient
from app.main import app
from app.models.document import SourceType, DocumentStatus
from app.services.document_service import DocumentService
from app.services.clause_service import ClauseService
from app.database.session import SessionLocal


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as test_client:
        yield test_client


@pytest.fixture
def db_session():
    from app.database.session import check_database_connection
    health = check_database_connection()
    if health["status"] == "disconnected":
        pytest.skip("PostgreSQL live instance is not running; skipping live DB test.")
    session = SessionLocal()
    try:
        yield session
    finally:
        session.close()


def test_analyze_plain_text(client: TestClient, db_session):
    """Test ingestion of plain text legal agreement."""
    sample_text = """
    TERMS OF SERVICE
    
    1. Introduction
    Welcome to our service. By accessing our platform, you agree to comply with these terms.
    
    2. Billing and Payments
    Your subscription automatically renews on a monthly basis until cancelled.
    All payments made are non-refundable unless required by applicable law.
    
    3. Limitation of Liability
    In no event shall the company be liable for any indirect, incidental, or consequential damages.
    """

    response = client.post(
        "/analyze/text",
        json={"text": sample_text, "title": "Test Terms of Service"},
    )
    assert response.status_code == 201
    data = response.json()
    assert "document_id" in data
    assert data["source_type"] == "text"
    assert data["total_clauses"] >= 3

    # Verify source references
    for clause in data["clauses"]:
        # Non-page documents must have NULL page_number
        assert clause["page_number"] is None
        assert clause["paragraph_number"] is not None
        assert isinstance(clause["clause_order"], int)
        assert len(clause["clause_text"]) > 0

    # Verify section detection
    sections = [c["section"] for c in data["clauses"] if c["section"]]
    assert any("Billing" in s or "Payments" in s for s in sections)

    # Verify persisted in PostgreSQL
    doc_id = data["document_id"]
    db_doc = DocumentService.get(db_session, doc_id)
    assert db_doc is not None
    assert db_doc.status == DocumentStatus.INDEXED
    clauses = ClauseService.list_by_document(db_session, doc_id)
    assert len(clauses) == data["total_clauses"]

    # Cleanup test document
    DocumentService.delete(db_session, doc_id)


def test_analyze_txt_file(client: TestClient, db_session):
    """Test ingestion of uploaded TXT file."""
    txt_content = (
        "PRIVACY POLICY\n\n"
        "Data Collection\n"
        "We collect personal information that you provide when registering an account.\n\n"
        "Data Sharing\n"
        "We do not sell your personal data to third parties without consent."
    ).encode("utf-8")

    response = client.post(
        "/analyze/file",
        files={"file": ("privacy_policy.txt", io.BytesIO(txt_content), "text/plain")},
    )
    assert response.status_code == 201
    data = response.json()
    assert data["source_type"] == "txt"
    assert data["filename"] == "privacy_policy.txt"
    assert data["total_clauses"] >= 2

    for clause in data["clauses"]:
        assert clause["page_number"] is None
        assert clause["paragraph_number"] is not None

    doc_id = data["document_id"]
    DocumentService.delete(db_session, doc_id)


def test_analyze_docx_file(client: TestClient, db_session):
    """Test ingestion of uploaded DOCX file created with python-docx."""
    doc = docx.Document()
    doc.add_heading("User Agreement", level=1)
    doc.add_paragraph("By accessing or using our application, you agree to these binding terms.")
    doc.add_heading("Dispute Resolution", level=2)
    doc.add_paragraph("Any disputes arising out of this agreement shall be settled through mandatory binding arbitration.")

    docx_io = io.BytesIO()
    doc.save(docx_io)
    docx_bytes = docx_io.getvalue()

    response = client.post(
        "/analyze/file",
        files={"file": ("agreement.docx", io.BytesIO(docx_bytes), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 201
    data = response.json()
    assert data["source_type"] == "docx"
    assert data["filename"] == "agreement.docx"
    assert data["total_clauses"] >= 2

    # Verify paragraph numbers are preserved and page numbers are None
    for clause in data["clauses"]:
        assert clause["page_number"] is None
        assert clause["paragraph_number"] is not None

    # Verify section was detected
    sections = [c["section"] for c in data["clauses"] if c["section"]]
    assert any("Dispute Resolution" in s for s in sections)

    doc_id = data["document_id"]
    DocumentService.delete(db_session, doc_id)


def test_analyze_pdf_file(client: TestClient, db_session):
    """Test ingestion of multi-page PDF created with PyMuPDF, verifying page numbers."""
    pdf_doc = fitz.open()

    # Page 1
    page1 = pdf_doc.new_page()
    page1.insert_text(
        (50, 72),
        "TERMS OF SERVICE\n\n1. General Terms\nUsers must be at least 18 years of age to register an account.",
        fontsize=12,
    )

    # Page 2
    page2 = pdf_doc.new_page()
    page2.insert_text(
        (50, 72),
        "2. Billing\nYour subscription automatically renews each month until you cancel.",
        fontsize=12,
    )

    pdf_bytes = pdf_doc.write()
    pdf_doc.close()

    response = client.post(
        "/analyze/file",
        files={"file": ("legal_notice.pdf", io.BytesIO(pdf_bytes), "application/pdf")},
    )
    assert response.status_code == 201
    data = response.json()
    assert data["source_type"] == "pdf"
    assert data["filename"] == "legal_notice.pdf"
    assert data["total_clauses"] >= 2

    # Verify page numbers are preserved (Page 1 and Page 2)
    page_numbers = {c["page_number"] for c in data["clauses"]}
    assert 1 in page_numbers
    assert 2 in page_numbers

    # Verify clause ordering
    orders = [c["clause_order"] for c in data["clauses"]]
    assert orders == list(range(len(orders)))

    # Verify section detection
    sections = [c["section"] for c in data["clauses"] if c["section"]]
    assert any("Billing" in s for s in sections)

    doc_id = data["document_id"]
    DocumentService.delete(db_session, doc_id)


def test_empty_plain_text(client: TestClient):
    """Verify empty plain text is rejected with 400."""
    response = client.post(
        "/analyze/text",
        json={"text": "   ", "title": "Empty Text"},
    )
    assert response.status_code == 400
    assert "empty" in response.json()["detail"].lower()


def test_empty_file_upload(client: TestClient):
    """Verify 0-byte file upload is rejected with 400."""
    response = client.post(
        "/analyze/file",
        files={"file": ("empty.txt", io.BytesIO(b""), "text/plain")},
    )
    assert response.status_code == 400
    assert "empty" in response.json()["detail"].lower()


def test_corrupted_pdf_upload(client: TestClient):
    """Verify corrupted PDF stream is rejected with 400."""
    corrupted_bytes = b"%PDF-1.4\nNOT_A_VALID_PDF_STREAM_CORRUPTED_12345"
    response = client.post(
        "/analyze/file",
        files={"file": ("corrupted.pdf", io.BytesIO(corrupted_bytes), "application/pdf")},
    )
    assert response.status_code == 400
    assert "corrupted" in response.json()["detail"].lower() or "invalid" in response.json()["detail"].lower()


def test_unsupported_file_type(client: TestClient):
    """Verify unsupported file extension is rejected with 400."""
    response = client.post(
        "/analyze/file",
        files={"file": ("script.py", io.BytesIO(b"print('hello')"), "text/x-python")},
    )
    assert response.status_code == 400
    assert "unsupported" in response.json()["detail"].lower()


def test_oversized_file(client: TestClient, monkeypatch):
    """Verify oversized file is rejected with 413."""
    import app.extraction.pipeline as pipeline_mod
    monkeypatch.setattr(pipeline_mod, "MAX_FILE_SIZE_BYTES", 100)

    oversized_bytes = b"A" * 200
    response = client.post(
        "/analyze/file",
        files={"file": ("oversized.txt", io.BytesIO(oversized_bytes), "text/plain")},
    )
    assert response.status_code == 413
    assert "exceeds" in response.json()["detail"].lower()
