import io
import zipfile
from typing import List, Tuple
import docx
from app.extraction.exceptions import CorruptedFileError, EmptyDocumentError, ExtractionError
from app.extraction.segmenter import ClauseSegmenter, ExtractedClause
from app.extraction.normalizer import clean_and_normalize_text


def extract_from_docx(file_bytes: bytes) -> Tuple[str, List[ExtractedClause]]:
    """
    Extracts text, preserves paragraph numbers, detects sections, and segments clauses from DOCX
    bytes using python-docx. Page numbers remain NULL.
    """
    if not file_bytes or len(file_bytes) == 0:
        raise EmptyDocumentError("The uploaded DOCX file is empty (0 bytes).")

    try:
        file_stream = io.BytesIO(file_bytes)
        doc = docx.Document(file_stream)
    except (zipfile.BadZipFile, docx.opc.exceptions.PackageNotFoundError) as e:
        raise CorruptedFileError(f"Corrupted or invalid DOCX document: {str(e)}")
    except Exception as e:
        raise ExtractionError(f"Failed to parse DOCX document: {str(e)}")

    raw_blocks: List[Tuple[str, None, int]] = []
    full_text_parts: List[str] = []
    para_counter = 1

    try:
        # Process document paragraphs
        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                # If paragraph style is an explicit Word Heading, ensure it has markdown-style hint
                if p.style and p.style.name and p.style.name.lower().startswith("heading"):
                    text = f"# {text}"
                raw_blocks.append((text, None, para_counter))
                full_text_parts.append(text)
                para_counter += 1

        # Process table text if present
        for table in doc.tables:
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    combined_row = " | ".join(row_texts)
                    raw_blocks.append((combined_row, None, para_counter))
                    full_text_parts.append(combined_row)
                    para_counter += 1
    except Exception as e:
        raise ExtractionError(f"Error reading DOCX paragraphs: {str(e)}")

    combined_text = "\n\n".join(full_text_parts)
    normalized_full_text = clean_and_normalize_text(combined_text)

    if not normalized_full_text:
        raise EmptyDocumentError("DOCX document contains no readable text.")

    segmenter = ClauseSegmenter()
    clauses = segmenter.process_blocks(raw_blocks)

    if not clauses:
        raise EmptyDocumentError("No readable clauses could be extracted from the DOCX file.")

    return normalized_full_text, clauses
